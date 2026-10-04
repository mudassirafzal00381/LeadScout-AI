"""LeadScout AI regression tests - no internet needed, takes about a minute.

    .venv\\Scripts\\python.exe tests\\run_tests.py

Covers: request parsing (Pakistan + international), phone formats per country,
reading phones/emails/social links from web pages, merging duplicates,
filters and scoring, Excel export, and the web app screens (with stand-in
search results, so no Google/OpenStreetMap traffic).
"""

import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)
os.environ["LEADSCOUT_OUTPUT_DIR"] = tempfile.mkdtemp(prefix="leadscout-test-")
os.environ["LEADSCOUT_LOG_DIR"] = tempfile.mkdtemp(prefix="leadscout-test-logs-")

from bs4 import BeautifulSoup as B  # noqa: E402

FAILED: list[str] = []


def check(condition, name):
    print(("PASS " if condition else "FAIL ") + name)
    if not condition:
        FAILED.append(name)


def section(title):
    print(f"\n--- {title}")


# --- Request parsing -----------------------------------------------------------
section("Request parsing")
from modules.nlp_parser import parse_filter_text, parse_user_request  # noqa: E402

for text, cat, loc, flt in [
    ("restaurants in Lahore that need a website", "restaurant", "Lahore", ["no_website"]),
    ("salons in Karachi", "salon", "Karachi", []),
    ("businesses in Lahore that need a website", "all", "Lahore", ["no_website"]),
    ("Find dental clinics near DHA Phase 5, Lahore with no website", "dentist", "DHA Phase 5, Lahore", ["no_website"]),
    ("hotels in London rated below 3.5", "hotel", "London", [("low_rating", 3.5)]),
    ("restaurants in New York that need a website", "restaurant", "New York", ["no_website"]),
    ("plumbers in London without a website", "plumber", "London", ["no_website"]),
    ("roofing companies in Houston, Texas", "roofer", "Houston, Texas", []),
    ("barbers in Toronto that need social media", "barber", "Toronto", ["no_social_media"]),
    ("nail salons in Sydney", "nail_salon", "Sydney", []),
    ("solar panel installers in Austin", "solar panel installer", "Austin", []),
    ("find me all local bakeries in Manchester", "bakery", "Manchester", []),
    ("car dealers in Los Angeles that need a website", "car_dealer", "Los Angeles", ["no_website"]),
]:
    r = parse_user_request(text)
    check((r["category"], r["location"], r["filters"]) == (cat, loc, flt), f"parse {text!r}")
check(parse_filter_text("needs a website, social media and phone number")[0]
      == ["no_website", "no_social_media", "no_phone"], "filter box: bare words")
check(parse_filter_text("no website, blah")[1] == ["blah"], "filter box: unknown words reported")

# --- Phone numbers ---------------------------------------------------------------
section("Phone numbers per country")
from modules.merge import normalize_phone, set_phone_region  # noqa: E402

for region, raw, want in [
    ("PK", "0321-4481300", "+923214481300"), ("PK", "(042) 111-444-411", "+9242111444411"),
    ("US", "(212) 555-0123", "+12125550123"), ("GB", "020 7946 0000", "+442079460000"),
    ("AE", "050 123 4567", "+971501234567"), ("CA", "(416) 555-0199", "+14165550199"),
    ("AU", "(02) 9374 4000", "+61293744000"), ("IN", "098765 43210", "+919876543210"),
    ("DE", "030 1234567", "+49301234567"), ("FR", "01 23 45 67 89", "+33123456789"),
    ("GB", "+1 416-555-0199", "+14165550199"), ("PK", "0092 300 1234567", "+923001234567"),
]:
    check(normalize_phone(raw, region) == want, f"[{region}] {raw} -> {want}")

# --- Web page reading --------------------------------------------------------------
section("Reading websites")
from modules.enrichment import _contact_page_url, extract_email, extract_phone, extract_social_links  # noqa: E402

for region, html, want in [
    ("US", "<p>Call us today: (312) 555-0147</p>", "+13125550147"),
    ("GB", "<footer>Tel: 0161 496 0000</footer>", "+441614960000"),
    ("PK", '<a href="tel:+92-42-35789823">Call</a>', "+924235789823"),
    ("US", "<p>Since 1998 - order #4581239087</p>", ""),
]:
    set_phone_region(region)
    check(extract_phone(B(html, "html.parser")) == want, f"phone on page [{region}] -> {want or 'none'}")
set_phone_region(None)
page = ('<a href="https://www.facebook.com/sharer.php?u=x">s</a><a href="https://m.facebook.com/Cafe/">f</a>'
        '<a href="https://instagram.com/cafe.pk">i</a><a href="https://[a-z]bad">bad</a>'
        '<a href="mailto:Orders@Cafe.pk">m</a><p>yourmail@gmail.com logo@2x.png</p>')
soup = B(page, "html.parser")
check(extract_social_links(soup, page, "https://cafe.pk") ==
      {"facebook": "https://www.facebook.com/Cafe", "instagram": "https://instagram.com/cafe.pk"},
      "social links (share buttons + malformed links ignored)")
check(extract_email(soup, "https://www.cafe.pk") == "orders@cafe.pk", "email (placeholders ignored)")
check(_contact_page_url(B('<a href="/contact-us">Contact</a>', "html.parser"), "https://cafe.pk")
      == "https://cafe.pk/contact-us", "contact page found")

# --- Merging, filters, scoring ---------------------------------------------------
section("Merging, filters and scoring")
from modules.filters import calculate_lead_score, filter_leads  # noqa: E402
from modules.lead import new_lead  # noqa: E402
from modules.merge import merge_and_deduplicate  # noqa: E402

set_phone_region("PK")
osm = [new_lead(name="Monal", address="Liberty Chowk", latitude=31.50985, longitude=74.34110, source="OpenStreetMap"),
       new_lead(name="KFC", latitude=31.47, longitude=74.27, source="OpenStreetMap")]
maps = [new_lead(name="Monal Lahore", phone="+92 42 35789823", website="themonal.com",
                 latitude=31.5097989, longitude=74.3410679, source="Google Maps"),
        new_lead(name="KFC", phone="042 111 532 532", latitude=31.52, longitude=74.35, source="Google Maps")]
merged = merge_and_deduplicate(osm, maps)
check(len(merged) == 3, "same business merged, chain branches kept apart")
monal = next(b for b in merged if "Monal" in b["name"])
check(monal["phone"] == "+924235789823" and monal["website"] == "https://themonal.com", "merged fields filled + normalized")
check(calculate_lead_score(new_lead(name="x")) == 100, "score: nothing = 100")
leads = [new_lead(name="A", phone="+921", email="a@a.pk"), new_lead(name="B", phone="+922"), new_lead(name="C")]
ranked = filter_leads(leads, ["has_phone"])
check([b["name"] for b in ranked] == ["A", "B"], "has_phone filter; equal scores: lead with email first")

# --- Excel export -----------------------------------------------------------------
section("Excel export")
from openpyxl import load_workbook  # noqa: E402

from modules.excel_export import export_to_excel  # noqa: E402

path = export_to_excel(ranked, category="salon", location="Lahore")
wb = load_workbook(path)
check(wb.sheetnames == ["Leads", "About"] and len(wb["About"]._images) == 1, "Excel: Leads + About sheet with logo")
check(wb["Leads"].freeze_panes == "B2" and wb["Leads"].auto_filter.ref, "Excel: frozen header + filters")

# --- Web app screens (stand-in search results, no internet) -------------------------
section("Web app")
from streamlit.runtime.context import ContextProxy  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402

import modules.maps_collector as mc  # noqa: E402
import modules.osm_collector as oc  # noqa: E402
import modules.pipeline as pl  # noqa: E402

ContextProxy.ip_address = property(lambda self: None)  # pretend: opened on this computer
oc.resolve_location = lambda loc: {"name": loc, "bbox": (0, 0, 1, 1), "lat": 0, "lon": 0, "country_code": "PK", "type": "city"}
pl.resolve_location = oc.resolve_location
pl.collect_from_osm = lambda c, l, m: [new_lead(name="Alpha Salon", category=c, latitude=31.5, longitude=74.3, source="OpenStreetMap"),
                                       new_lead(name="Beta Salon", category=c, latitude=31.6, longitude=74.3, source="OpenStreetMap")]
mc.collect_from_google_maps = lambda *a, **k: [new_lead(name="Gamma Salon", category="Beauty salon", phone="+92 300 1112233",
                                                        latitude=31.9, longitude=74.3, source="Google Maps")]


def fake_lookup(leads, location, cap):
    for lead in leads:
        if lead["name"] == "Beta Salon":
            lead["phone"] = "+923004445566"
    return {"looked_up": len(leads), "phones_found": 1}


mc.lookup_missing_details = fake_lookup
pl.enrich_with_social_and_contact = lambda leads: leads
pl.time.sleep = lambda s: None


def app():
    return AppTest.from_file(str(ROOT / "app.py"), default_timeout=90)


at = app().run()
check(not at.exception and any("looking for" in t.label for t in at.text_input), "app opens locally without login")
at.text_input[0].input("salons in Lahore that need a website").run()
[b for b in at.button if "Find Leads" in b.label][0].click().run()
names = set(at.dataframe[0].value["Business Name"]) if at.dataframe else set()
check(names == {"Beta Salon", "Gamma Salon"}, "only leads with a phone number are shown")
check(any("left out" in i.value for i in at.info), "results say how many were left out")
check(not at.exception, "search completes without errors")

ContextProxy.ip_address = property(lambda self: "203.0.113.9")  # another device
at = app().run()
check(any("no password set" in e.value for e in at.error), "other devices blocked when no password")
os.environ["LEADSCOUT_PASSWORD"] = "T3st-pass!"
at = app().run()
check(not any("looking for" in t.label for t in at.text_input), "password page shown")
at.text_input[0].input("T3st-pass!").run()
at.button[0].click().run()
check(any("looking for" in t.label for t in at.text_input), "correct password opens the app")
del os.environ["LEADSCOUT_PASSWORD"]

print(f"\n{'ALL TESTS PASSED' if not FAILED else f'{len(FAILED)} FAILED: ' + '; '.join(FAILED)}")
sys.exit(1 if FAILED else 0)
