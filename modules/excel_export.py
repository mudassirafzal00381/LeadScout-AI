"""Export leads to a clean, formatted Excel file (openpyxl)."""

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config

# --- Styles ------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MISSING_RED = PatternFill("solid", start_color="FFC7CE")     # key gap: website / phone
MISSING_ORANGE = PatternFill("solid", start_color="FCE4D6")  # secondary gap
LINK_FONT = Font(color="0563C1", underline="single")
THIN_BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))

MIN_WIDTH, MAX_WIDTH = 8, 50

OTHER_SOCIAL = ("twitter", "tiktok", "youtube", "whatsapp")
PLATFORM_LABELS = {"twitter": "X/Twitter", "tiktok": "TikTok", "youtube": "YouTube",
                   "whatsapp": "WhatsApp"}


def _address(b: dict) -> str:
    if b.get("address"):
        return b["address"]
    if b.get("latitude") is not None and b.get("longitude") is not None:
        return f"{b['latitude']:.6f}, {b['longitude']:.6f}"
    return ""


def _social(platform: str):
    return lambda b: (b.get("social_media") or {}).get(platform, "")


def _other_social(b: dict) -> str:
    social = b.get("social_media") or {}
    return "\n".join(f"{PLATFORM_LABELS[p]}: {social[p]}" for p in OTHER_SOCIAL if social.get(p))


# (header, value getter, highlight fill when empty, is_link)
COLUMNS = [
    ("Business Name", lambda b: b.get("name", ""), None, False),
    ("Category", lambda b: b.get("category", ""), None, False),
    ("Contact Number", lambda b: b.get("phone", ""), MISSING_RED, False),
    ("Email", lambda b: b.get("email", ""), MISSING_ORANGE, False),
    ("Website", lambda b: b.get("website", ""), MISSING_RED, True),
    ("Facebook", _social("facebook"), MISSING_ORANGE, True),
    ("Instagram", _social("instagram"), MISSING_ORANGE, True),
    ("LinkedIn", _social("linkedin"), MISSING_ORANGE, True),
    ("Other Social", _other_social, None, False),
    ("Address/Location", _address, MISSING_ORANGE, False),
    ("Rating", lambda b: b.get("rating"), None, False),
    ("Lead Score", lambda b: b.get("lead_score"), None, False),
    ("Missing Fields", lambda b: ", ".join(b.get("missing_fields") or []), None, False),
    # Extra columns: where the lead came from, with a link to the listing.
    ("Source", lambda b: b.get("source", ""), None, False),
    ("Listing Link", lambda b: (b.get("source_url") or "").split(" ; ")[0], None, True),
]


# --- Helpers -----------------------------------------------------------------

def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (text or "").lower()).strip("_") or "all"


def build_filename(category: str, location: str, when: datetime | None = None) -> str:
    """e.g. build_filename("restaurant", "Lahore, Pakistan") -> leads_restaurant_lahore_pakistan_2026-09-26.xlsx"""
    date = (when or datetime.now()).strftime("%Y-%m-%d")
    return f"leads_{_slug(category)}_{_slug(location)}_{date}.xlsx"


def _clean(value):
    """Make a value safe for Excel: strip illegal control characters."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value).strip()
    return value


def _resolve_path(filename: str) -> Path:
    """Save inside config.OUTPUT_DIR (unless a directory is given) without overwriting."""
    path = Path(filename)
    if path.suffix.lower() != ".xlsx":
        path = path.with_name(path.name + ".xlsx")
    if not path.is_absolute() and path.parent == Path("."):
        path = config.OUTPUT_DIR / path
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():  # keep earlier exports from the same day
        path = path.with_name(f"{path.stem}_{datetime.now():%H%M%S}{path.suffix}")
    return path


def _add_about_sheet(wb: Workbook, category: str, location: str) -> None:
    """Second sheet: search details, data sources and the required OSM attribution."""
    ws = wb.create_sheet("About")
    rows = [
        ("LeadScout AI - lead export", None),
        ("", None),
        ("Search", f"{category.replace('_', ' ') or 'all'} in {location or '-'}"),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", None),
        ("Data sources", None),
        ("OpenStreetMap", "Map data © OpenStreetMap contributors, available under the "
                          "Open Database License (ODbL): https://www.openstreetmap.org/copyright"),
        ("Google Maps", "Public business listings, collected in small numbers for reference."),
        ("Websites", "Social media links and emails found on each business's own homepage."),
        ("", None),
        ("Colour legend", None),
        ("Red cell", "Key gap: no phone number, or no real website (none, or only a social page)."),
        ("Orange cell", "Other missing info (email, social profiles, address) or a broken website."),
        ("Lead Score", "0-100. Higher = more gaps = more services you could offer."),
    ]
    for label, value in rows:
        ws.append([label, value])
    ws["A1"].font = Font(bold=True, size=14)
    for cell in ("A6", "A11"):
        ws[cell].font = Font(bold=True)
    ws["A12"].fill, ws["A13"].fill = MISSING_RED, MISSING_ORANGE
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


def _sort_key(b: dict):
    score = b.get("lead_score")
    return (score is not None, score or 0)


# --- Public API --------------------------------------------------------------

COLUMN_HEADERS = [header for header, *_ in COLUMNS]


def table_rows(business_list: list[dict]) -> list[dict]:
    """Leads as rows with the same columns/values as the Excel export (for on-screen tables)."""
    rows = []
    for business in sorted(business_list, key=_sort_key, reverse=True):
        row = {}
        for header, getter, _, _ in COLUMNS:
            try:
                value = _clean(getter(business))
            except Exception:  # noqa: BLE001
                value = ""
            row[header] = value
        rows.append(row)
    return rows


def export_to_excel(business_list: list[dict], filename: str | None = None, *,
                    category: str = "", location: str = "") -> Path:
    """Write leads to a formatted .xlsx file and return its path.

    filename: a name (saved in config.OUTPUT_DIR) or a full path. If omitted,
              it is built as leads_{category}_{location}_{date}.xlsx.
    Rows are sorted by Lead Score, highest first.
    """
    filename = filename or build_filename(category, location)
    rows = sorted(business_list, key=_sort_key, reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    # Header
    ws.append([header for header, *_ in COLUMNS])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    widths = [len(header) + 4 for header, *_ in COLUMNS]  # room for the filter arrow

    for r, business in enumerate(rows, start=2):
        missing = business.get("missing_fields") or []
        for c, (header, getter, missing_fill, is_link) in enumerate(COLUMNS, start=1):
            try:
                value = _clean(getter(business))
            except Exception:  # noqa: BLE001 - one odd record shouldn't break the export
                value = ""
            cell = ws.cell(row=r, column=c)
            cell.value = value if value != "" else None
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # never let scraped text run as a formula
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=header == "Other Social")

            if value in ("", None):
                if missing_fill is not None:
                    cell.fill = missing_fill
            elif is_link and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.font = LINK_FONT

            # A social page used as the "website", or a broken site, is still a gap.
            if header == "Website" and value:
                if "website" in missing:
                    cell.fill = MISSING_RED
                elif "website (broken)" in missing:
                    cell.fill = MISSING_ORANGE

            if header == "Rating" and isinstance(value, float):
                cell.number_format = "0.0"
            if header in ("Rating", "Lead Score"):
                cell.alignment = Alignment(horizontal="center", vertical="top")

            longest_line = max((len(line) for line in str(value or "").split("\n")), default=0)
            widths[c - 1] = max(widths[c - 1], longest_line + 2)

    # Column widths
    for c, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(MIN_WIDTH, min(width, MAX_WIDTH))

    last_row = max(ws.max_row, 2)
    last_col = get_column_letter(len(COLUMNS))

    # Freeze header row + name column; add filter dropdowns to the header.
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    # Lead Score colour scale: white (0) -> orange -> red (100) = hottest.
    score_col = get_column_letter([h for h, *_ in COLUMNS].index("Lead Score") + 1)
    ws.conditional_formatting.add(
        f"{score_col}2:{score_col}{last_row}",
        ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                       mid_type="num", mid_value=50, mid_color="FFD966",
                       end_type="num", end_value=100, end_color="F8696B"),
    )

    _add_about_sheet(wb, category, location)

    path = _resolve_path(filename)
    try:
        wb.save(path)
    except PermissionError:  # usually: the file is open in Excel
        path = path.with_name(f"{path.stem}_{datetime.now():%H%M%S}{path.suffix}")
        wb.save(path)
    return path
